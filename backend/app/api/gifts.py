"""礼品单：订单台账管理（日期/下单时间/店铺/关键词/规格/金额/佣金/旺旺号/订单号/评论/结款）。"""

from __future__ import annotations

import secrets
import uuid
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel

from backend.app.api.auth import get_current_user, visible_store_ids
from backend.app.core.db import DB_PATH, get_db
from backend.app.core.logs import log_op

router = APIRouter()

REVIEW_STATUSES = ("none", "reviewed")
SETTLE_STATUSES = ("unsettled", "settled")
REVIEW_LABELS = {"none": "未评论", "reviewed": "已评论"}
SETTLE_LABELS = {"unsettled": "未结款", "settled": "已结款"}


class GiftIn(BaseModel):
    order_no: str = ""
    store_id: int = 0
    store_name: str = ""
    keyword: str = ""
    spec: str = ""
    price: float = 0
    commission: float = 0
    wangwang: str = ""
    order_time: str = ""
    review_status: str = "none"
    settle_status: str = "unsettled"
    image: str = ""


class GiftStatusIn(BaseModel):
    status: str


class GiftBatchIn(BaseModel):
    ids: list[int]
    review_status: str | None = None
    settle_status: str | None = None


class GiftBatchDeleteIn(BaseModel):
    ids: list[int]


class GiftBatchCreateIn(BaseModel):
    date: str = ""
    start_time: str = ""
    store_id: int = 0
    store_name: str = ""
    keyword: str = ""
    spec: str = ""
    price: float = 0
    commission: float = 0
    quantity: int = 1
    image: str = ""


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _store_clause(user: dict, alias: str = "g") -> tuple[str, list]:
    """SaaS 隔离：member 只能看到/操作绑定店铺的礼品单；超管/管理员不限。"""
    visible = visible_store_ids(user)
    if visible is None:
        return "", []
    col = f"{alias}.store_id" if alias else "store_id"
    if not visible:
        return f" AND {col} IN (0)", []
    ids = ",".join(str(i) for i in visible)
    return f" AND {col} IN ({ids})", []


def _ensure_gift_access(user: dict, row) -> None:
    """校验单条礼品单是否属于用户可见店铺。"""
    visible = visible_store_ids(user)
    if visible is None:
        return
    if row["store_id"] not in visible:
        raise HTTPException(status_code=403, detail="无权操作该礼品单")


def _get_gift_or_404(db, gift_id: int):
    row = db.execute(
        """
        SELECT g.*, COALESCE(s.name, '') AS linked_store_name
        FROM gifts g
        LEFT JOIN stores s ON s.id = g.store_id
        WHERE g.id = ?
        """,
        (gift_id,),
    ).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="礼品单不存在")
    return row


def _payload(row) -> dict:
    return {
        "id": row["id"],
        "store_id": row["store_id"],
        "store_name": (
            (row["linked_store_name"] or "") if row["store_id"] != 0 else (row["store_name"] or "")
        )
        or "未关联店铺",
        "order_no": row["order_no"],
        "keyword": row["keyword"],
        "spec": row["spec"],
        "price": row["price"],
        "commission": row["commission"],
        "wangwang": row["wangwang"],
        "order_time": row["order_time"],
        "review_status": row["review_status"],
        "settle_status": row["settle_status"],
        "status": row["status"],
        "recipient": row["recipient"],
        "gift_name": row["gift_name"],
        "quantity": row["quantity"],
        "image": row["image"] or "",
        "created_at": row["created_at"],
    }


def _validate(body: GiftIn, has_image: bool = False, existing_store_name: str = "") -> dict:
    keyword = body.keyword.strip()
    spec = body.spec.strip()
    wangwang = body.wangwang.strip()
    if body.store_id == 0 and not body.store_name.strip() and not existing_store_name:
        raise HTTPException(status_code=400, detail="请选择或输入店铺")
    if not keyword and not has_image:
        raise HTTPException(status_code=400, detail="关键词不能为空（可填文字或粘贴图片）")
    if body.price <= 0:
        raise HTTPException(status_code=400, detail="金额必须大于 0")
    if not (0 <= body.price <= 999999):
        raise HTTPException(status_code=400, detail="金额超出范围")
    if not (0 <= body.commission <= 999999):
        raise HTTPException(status_code=400, detail="佣金超出范围")
    if body.review_status not in REVIEW_STATUSES:
        raise HTTPException(status_code=400, detail="评论状态不正确")
    if body.settle_status not in SETTLE_STATUSES:
        raise HTTPException(status_code=400, detail="结款状态不正确")
    return {"keyword": keyword, "spec": spec, "wangwang": wangwang}


def _resolve_order_no(db, body: GiftIn, exclude_id: int | None = None) -> str:
    order_no = body.order_no.strip()
    if order_no:
        if not order_no.isdigit():
            raise HTTPException(status_code=400, detail="订单编号只能是数字")
        if len(order_no) > 40:
            raise HTTPException(status_code=400, detail="订单号过长（最多 40 个字符）")
        if exclude_id is not None:
            exists = db.execute(
                "SELECT id FROM gifts WHERE order_no = ? AND id != ?",
                (order_no, exclude_id),
            ).fetchone()
        else:
            exists = db.execute("SELECT id FROM gifts WHERE order_no = ?", (order_no,)).fetchone()
        if exists:
            raise HTTPException(status_code=400, detail=f"订单号「{order_no}」已存在，请检查后重试")
        return order_no
    return ""


@router.get("")
def list_gifts(
    store_id: int | None = None,
    review_status: str | None = None,
    settle_status: str | None = None,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    query = """
        SELECT g.*, COALESCE(s.name, '') AS linked_store_name
        FROM gifts g
        LEFT JOIN stores s ON s.id = g.store_id
        WHERE 1 = 1
    """
    params: list = []
    sf, sp = _store_clause(user)
    query += sf
    params += sp
    if store_id is not None:
        query += " AND g.store_id = ?"
        params.append(store_id)
    if review_status:
        query += " AND g.review_status = ?"
        params.append(review_status)
    if settle_status:
        query += " AND g.settle_status = ?"
        params.append(settle_status)
    query += " ORDER BY COALESCE(g.order_time, g.created_at) ASC, g.id ASC"
    rows = db.execute(query, params).fetchall()
    return {"items": [_payload(row) for row in rows]}


@router.get("/export")
def export_gifts(
    keyword: str = "",
    store_id: int | None = None,
    store_name: str = "",
    review_status: str | None = None,
    settle_status: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    ids: str = "",
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> StreamingResponse:
    query = """
        SELECT g.*, COALESCE(s.name, '') AS linked_store_name
        FROM gifts g
        LEFT JOIN stores s ON s.id = g.store_id
        WHERE 1 = 1
    """
    params: list = []
    sf, sp = _store_clause(user)
    query += sf
    params += sp
    if store_id is not None:
        query += " AND g.store_id = ?"
        params.append(store_id)
    if review_status:
        query += " AND g.review_status = ?"
        params.append(review_status)
    if settle_status:
        query += " AND g.settle_status = ?"
        params.append(settle_status)
    if store_name.strip():
        query += " AND COALESCE(CASE WHEN g.store_id != 0 THEN s.name ELSE g.store_name END, '') = ?"
        params.append(store_name.strip())
    if keyword.strip():
        kw = f"%{keyword.strip()}%"
        query += " AND (g.order_no LIKE ? OR g.wangwang LIKE ? OR g.keyword LIKE ? OR g.spec LIKE ?)"
        params += [kw, kw, kw, kw]
    if date_from:
        query += " AND COALESCE(g.order_time, g.created_at) >= ?"
        params.append(date_from)
    if date_to:
        query += " AND COALESCE(g.order_time, g.created_at) <= ?"
        params.append(date_to)
    if ids.strip():
        id_list = [int(x) for x in ids.split(",") if x.strip().isdigit()]
        if id_list:
            placeholders = ",".join("?" for _ in id_list)
            query += f" AND g.id IN ({placeholders})"
            params.extend(id_list)
    query += " ORDER BY COALESCE(g.order_time, g.created_at) ASC, g.id ASC"
    rows = db.execute(query, params).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "礼品单"
    ws.append(["日期", "下单时间", "店铺", "关键词", "规格", "金额", "佣金", "旺旺号", "订单编号", "评论状态", "结款状态"])
    for row in rows:
        ot = row["order_time"] or row["created_at"] or ""
        ws.append(
            [
                ot[:10] if ot else "",
                (ot.split("T")[-1].split(" ")[-1][:5]) if ot else "",
                (row["linked_store_name"] or "") if row["store_id"] != 0 else (row["store_name"] or ""),
                row["keyword"],
                row["spec"],
                row["price"],
                row["commission"],
                row["wangwang"],
                row["order_no"],
                REVIEW_LABELS.get(row["review_status"], row["review_status"]),
                SETTLE_LABELS.get(row["settle_status"], row["settle_status"]),
            ]
        )
    widths = [12, 18, 20, 18, 14, 10, 10, 16, 24, 10, 10]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"礼品单_{datetime.now().strftime('%Y%m%d%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/image-upload")
async def upload_gift_image_standalone(
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
) -> dict:
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="请选择图片")
    ext = Path(file.filename).suffix.lower()
    if ext not in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        raise HTTPException(status_code=400, detail="仅支持 PNG/JPG/GIF/WebP 图片")
    content = await file.read()
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="图片不能超过 2MB")
    images_dir = DB_PATH.parent / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    name = f"u{uuid.uuid4().hex[:12]}{ext}"
    (images_dir / name).write_bytes(content)
    return {"url": f"/api/images/{name}"}


@router.post("/batch-create")
def batch_create_gifts(
    body: GiftBatchCreateIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    quantity = body.quantity
    if not (1 <= quantity <= 100):
        raise HTTPException(status_code=400, detail="下单数量需在 1-100 之间")
    keyword = body.keyword.strip()
    image = body.image.strip()
    if not keyword and not image:
        raise HTTPException(status_code=400, detail="关键词不能为空（可填文字或粘贴图片）")
    if body.price <= 0:
        raise HTTPException(status_code=400, detail="金额必须大于 0")
    if not (0 <= body.commission <= 999999):
        raise HTTPException(status_code=400, detail="佣金超出范围")
    if body.store_id != 0:
        store = db.execute("SELECT id FROM stores WHERE id = ?", (body.store_id,)).fetchone()
        if not store:
            raise HTTPException(status_code=400, detail="所选店铺不存在")
        store_id = body.store_id
        gift_store_name = ""
    else:
        gift_store_name = body.store_name.strip()
        if not gift_store_name:
            raise HTTPException(status_code=400, detail="请选择或输入店铺")
        row = db.execute("SELECT id FROM stores WHERE name = ?", (gift_store_name,)).fetchone()
        if row:
            store_id = row["id"]
            gift_store_name = ""
        else:
            store_id = 0

    visible = visible_store_ids(user)
    if visible is not None and store_id not in visible:
        raise HTTPException(status_code=403, detail="只能创建自己绑定店铺的礼品单")

    now = datetime.now()
    base_date = now.date()
    if body.date:
        try:
            base_date = datetime.strptime(body.date, "%Y-%m-%d").date()
        except ValueError:
            raise HTTPException(status_code=400, detail="下单日期格式不正确")
    base = datetime.combine(base_date, now.time())
    if body.start_time.strip():
        try:
            start_h, start_m = (int(x) for x in body.start_time.strip().split(":"))
            base = base.replace(hour=start_h, minute=start_m, second=0, microsecond=0)
        except (ValueError, TypeError):
            raise HTTPException(status_code=400, detail="开始时间格式不正确（应为 HH:mm）")
    created_at = _now()
    items = []
    for i in range(quantity):
        ts = base + timedelta(minutes=15 * i)
        order_time = ts.strftime("%Y-%m-%d %H:%M:%S")
        cur = db.execute(
            """
            INSERT INTO gifts (store_id, store_name, order_no, keyword, spec, price, commission, wangwang, order_time,
                               review_status, settle_status, status, created_at, image)
            VALUES (?, ?, ?, ?, ?, ?, ?, '', ?, 'none', 'unsettled', 'pending', ?, ?)
            """,
            (
                store_id,
                gift_store_name,
                "",
                keyword,
                body.spec.strip(),
                body.price,
                body.commission,
                order_time,
                created_at,
                image,
            ),
        )
        items.append(_payload(_get_gift_or_404(db, cur.lastrowid)))
    log_op(db, user, "gifts", "create", keyword or image, f"批量生成 {quantity} 条礼品单")
    return {"items": items}


@router.post("/batch")
def batch_update_gifts(
    body: GiftBatchIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    if not body.ids:
        raise HTTPException(status_code=400, detail="请先选择要修改的礼品单")
    if body.review_status is None and body.settle_status is None:
        raise HTTPException(status_code=400, detail="请选择要修改的评论状态或结款状态")
    if body.review_status is not None and body.review_status not in REVIEW_STATUSES:
        raise HTTPException(status_code=400, detail="评论状态不正确")
    if body.settle_status is not None and body.settle_status not in SETTLE_STATUSES:
        raise HTTPException(status_code=400, detail="结款状态不正确")
    placeholders = ",".join("?" for _ in body.ids)
    scf, _ = _store_clause(user, alias="")
    if body.review_status is not None:
        db.execute(
            f"UPDATE gifts SET review_status = ? WHERE id IN ({placeholders})" + scf,
            (body.review_status, *body.ids),
        )
    if body.settle_status is not None:
        db.execute(
            f"UPDATE gifts SET settle_status = ? WHERE id IN ({placeholders})" + scf,
            (body.settle_status, *body.ids),
        )
    log_op(db, user, "gifts", "batch", "", f"批量更新 {len(body.ids)} 单")
    return {"ok": True, "count": len(body.ids)}


@router.post("/batch-delete")
def batch_delete_gifts(
    body: GiftBatchDeleteIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    if not body.ids:
        raise HTTPException(status_code=400, detail="请先选择要删除的礼品单")
    placeholders = ",".join("?" for _ in body.ids)
    scf, _ = _store_clause(user, alias="")
    db.execute(f"DELETE FROM gifts WHERE id IN ({placeholders})" + scf, body.ids)
    log_op(db, user, "gifts", "delete", "", f"批量删除 {len(body.ids)} 单")
    return {"ok": True, "count": len(body.ids)}


@router.post("")
def create_gift(
    body: GiftIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    fields = _validate(body, has_image=bool(body.image.strip()))
    if body.store_id != 0:
        store = db.execute("SELECT id FROM stores WHERE id = ?", (body.store_id,)).fetchone()
        if not store:
            raise HTTPException(status_code=400, detail="所选店铺不存在")
        store_id = body.store_id
        gift_store_name = ""
    else:
        gift_store_name = body.store_name.strip()
        if not gift_store_name:
            raise HTTPException(status_code=400, detail="请选择或输入店铺")
        store_id = 0
    visible = visible_store_ids(user)
    if visible is not None and store_id not in visible:
        raise HTTPException(status_code=403, detail="只能创建自己绑定店铺的礼品单")
    order_no = _resolve_order_no(db, body)
    order_time = body.order_time.strip() or _now()
    cur = db.execute(
        """
        INSERT INTO gifts (store_id, store_name, order_no, keyword, spec, price, commission, wangwang, order_time,
                           review_status, settle_status, status, created_at, image)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
        """,
        (
            store_id,
            gift_store_name,
            order_no,
            fields["keyword"],
            fields["spec"],
            body.price,
            body.commission,
            fields["wangwang"],
            order_time,
            body.review_status,
            body.settle_status,
            _now(),
            body.image.strip(),
        ),
    )
    item = _payload(_get_gift_or_404(db, cur.lastrowid))
    log_op(db, user, "gifts", "create", order_no, f"新增礼品单 {order_no}")
    return {"item": item}


@router.put("/{gift_id}")
def update_gift(
    gift_id: int,
    body: GiftIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    row = _get_gift_or_404(db, gift_id)
    _ensure_gift_access(user, row)
    fields = _validate(
        body,
        has_image=bool(body.image.strip() or row["image"]),
        existing_store_name=row["store_name"] or "",
    )
    if body.store_id != 0:
        store = db.execute("SELECT id FROM stores WHERE id = ?", (body.store_id,)).fetchone()
        if not store:
            raise HTTPException(status_code=400, detail="所选店铺不存在")
        store_id = body.store_id
        gift_store_name = ""
    elif body.store_name.strip():
        store_id = 0
        gift_store_name = body.store_name.strip()
    else:
        store_id = row["store_id"]
        gift_store_name = row["store_name"] or ""
    order_no = body.order_no.strip()
    if order_no:
        if order_no != row["order_no"] and not order_no.isdigit():
            raise HTTPException(status_code=400, detail="订单编号只能是数字")
        if len(order_no) > 40:
            raise HTTPException(status_code=400, detail="订单号过长（最多 40 个字符）")
        exists = db.execute(
            "SELECT id FROM gifts WHERE order_no = ? AND id != ?",
            (order_no, gift_id),
        ).fetchone()
        if exists:
            raise HTTPException(status_code=400, detail=f"订单号「{order_no}」已存在，请检查后重试")
    order_time = body.order_time.strip() or row["order_time"] or row["created_at"]
    db.execute(
        """
        UPDATE gifts
        SET store_id = ?, store_name = ?, order_no = ?, keyword = ?, spec = ?, price = ?, commission = ?,
            wangwang = ?, order_time = ?, review_status = ?, settle_status = ?
        WHERE id = ?
        """,
        (
            store_id,
            gift_store_name,
            order_no,
            fields["keyword"],
            fields["spec"],
            body.price,
            body.commission,
            fields["wangwang"],
            order_time,
            body.review_status,
            body.settle_status,
            gift_id,
        ),
    )
    item = _payload(_get_gift_or_404(db, gift_id))
    log_op(db, user, "gifts", "update", order_no, f"编辑礼品单 {row['order_no']}")
    return {"item": item}


@router.post("/{gift_id}/review")
def update_gift_review(
    gift_id: int,
    body: GiftStatusIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    if body.status not in REVIEW_STATUSES:
        raise HTTPException(status_code=400, detail="评论状态不正确")
    row = _get_gift_or_404(db, gift_id)
    _ensure_gift_access(user, row)
    db.execute("UPDATE gifts SET review_status = ? WHERE id = ?", (body.status, gift_id))
    item = _payload(_get_gift_or_404(db, gift_id))
    log_op(db, user, "gifts", "review", row["order_no"], f"评论状态改为「{REVIEW_LABELS[body.status]}」")
    return {"item": item}


@router.post("/{gift_id}/settle")
def update_gift_settle(
    gift_id: int,
    body: GiftStatusIn,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    if body.status not in SETTLE_STATUSES:
        raise HTTPException(status_code=400, detail="结款状态不正确")
    row = _get_gift_or_404(db, gift_id)
    _ensure_gift_access(user, row)
    db.execute("UPDATE gifts SET settle_status = ? WHERE id = ?", (body.status, gift_id))
    item = _payload(_get_gift_or_404(db, gift_id))
    log_op(db, user, "gifts", "settle", row["order_no"], f"结款状态改为「{SETTLE_LABELS[body.status]}」")
    return {"item": item}


@router.post("/{gift_id}/image")
async def upload_gift_image(
    gift_id: int,
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    row = _get_gift_or_404(db, gift_id)
    _ensure_gift_access(user, row)
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="请选择二维码图片")
    ext = Path(file.filename).suffix.lower()
    if ext not in (".png", ".jpg", ".jpeg", ".gif", ".webp"):
        raise HTTPException(status_code=400, detail="仅支持 PNG/JPG/GIF/WebP 图片")
    content = await file.read()
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="图片不能超过 2MB")
    qr_dir = DB_PATH.parent / "images"
    qr_dir.mkdir(parents=True, exist_ok=True)
    name = f"g{gift_id}_{uuid.uuid4().hex[:12]}{ext}"
    (qr_dir / name).write_bytes(content)
    old = row["image"]
    if old and old.startswith("/api/images/"):
        old_path = qr_dir / Path(old.rsplit("/", 1)[-1])
        if old_path.exists() and old_path.name != name:
            try:
                old_path.unlink()
            except OSError:
                pass
    db.execute("UPDATE gifts SET image = ? WHERE id = ?", (f"/api/images/{name}", gift_id))
    item = _payload(_get_gift_or_404(db, gift_id))
    log_op(db, user, "gifts", "image", row["order_no"], "上传图片")
    return {"item": item}


@router.post("/{gift_id}/image/clear")
def clear_gift_image(
    gift_id: int,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    row = _get_gift_or_404(db, gift_id)
    _ensure_gift_access(user, row)
    old = row["image"]
    if old and old.startswith("/api/images/"):
        old_path = DB_PATH.parent / "images" / Path(old.rsplit("/", 1)[-1])
        if old_path.exists():
            try:
                old_path.unlink()
            except OSError:
                pass
    db.execute("UPDATE gifts SET image = '' WHERE id = ?", (gift_id,))
    item = _payload(_get_gift_or_404(db, gift_id))
    log_op(db, user, "gifts", "image", row["order_no"], "移除图片")
    return {"item": item}


@router.delete("/{gift_id}")
def delete_gift(
    gift_id: int,
    user: dict = Depends(get_current_user),
    db=Depends(get_db),
) -> dict:
    row = _get_gift_or_404(db, gift_id)
    _ensure_gift_access(user, row)
    db.execute("DELETE FROM gifts WHERE id = ?", (gift_id,))
    log_op(db, user, "gifts", "delete", row["order_no"], f"删除礼品单 {row['order_no']}")
    return {"ok": True}
